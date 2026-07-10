from pathlib import Path

import openpyxl
import pymupdf
from docx import Document as DocxFile

SUPPORTED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/plain",
}


def extract_text(file_path: str, mime_type: str) -> str:

    if mime_type == "application/pdf":
        return _extract_pdf(file_path)

    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(file_path)

    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _extract_xlsx(file_path)

    if mime_type == "text/plain":
        return _extract_txt(file_path)

    raise ValueError(f"Unsupported mime type: {mime_type}")


def _extract_pdf(file_path: str) -> str:

    with pymupdf.open(file_path) as document:
        return "\n".join(page.get_text() for page in document)


def _extract_docx(file_path: str) -> str:

    document = DocxFile(file_path)

    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_xlsx(file_path: str) -> str:

    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

    sheets = []

    for sheet in workbook.worksheets:

        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell.strip() for cell in row)]

        if not rows:
            continue

        header, *body = rows

        lines = [
            f"## {sheet.title}",
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
            *(f"| {' | '.join(row)} |" for row in body),
        ]

        sheets.append("\n".join(lines))

    return "\n\n".join(sheets)


def _extract_txt(file_path: str) -> str:

    return Path(file_path).read_text(encoding="utf-8")
